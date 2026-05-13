"""
utils/excel_handler.py
-----------------------
Responsavel por toda a interacao com o arquivo Excel de saida.

Colunas (ordem definida pelo desafio):
  1. title               - Titulo da noticia
  2. date                - Data de publicacao
  3. description         - Descricao/subtitulo (se disponivel)
  4. image_filename      - Nome do arquivo de imagem baixado
  5. search_phrase_count - Ocorrencias da frase de pesquisa (titulo+descricao)
  6. contains_money      - True/False se ha mencao a valor monetario
"""

import logging
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("nytimes_scraper.excel_handler")

COLUMNS = [
    "title",
    "date",
    "description",
    "image_filename",
    "search_phrase_count",
    "contains_money",
]

COLUMN_WIDTHS = {
    "title":               55,
    "date":                15,
    "description":         75,
    "image_filename":      45,
    "search_phrase_count": 22,
    "contains_money":      16,
}

_HEADER_FONT      = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL      = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_ALIGN     = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN_BORDER      = Border(bottom=Side(style="thin", color="D3D3D3"))
_ROW_FILL_EVEN    = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")


class ExcelHandler:
    """
    Gerencia a criacao, escrita e salvamento do arquivo Excel de resultados.

    Uso:
        handler = ExcelHandler(config)
        handler.initialize()
        handler.append_row(data_dict)
        handler.save()
    """

    def __init__(self, config):
        self.config      = config
        self.output_path = Path(config.output.excel_folder) / config.output.excel_filename
        self._workbook   = None
        self._worksheet  = None
        self._row_index  = 2

    def initialize(self) -> None:
        """Cria o workbook e escreve o cabecalho formatado."""
        self._workbook  = Workbook()
        self._worksheet = self._workbook.active
        self._worksheet.title = "News"
        self._worksheet.freeze_panes = "A2"

        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell           = self._worksheet.cell(row=1, column=col_idx, value=col_name)
            cell.font      = _HEADER_FONT
            cell.fill      = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        self._worksheet.row_dimensions[1].height = 30
        self._row_index = 2
        logger.info(f"Excel inicializado. Destino: {self.output_path}")

    def append_row(self, data: dict) -> None:
        """Adiciona uma linha de dados ao Excel."""
        if self._worksheet is None:
            self.initialize()

        is_even = (self._row_index % 2 == 0)

        for col_idx, col_name in enumerate(COLUMNS, start=1):
            value = data.get(col_name, "")

            if isinstance(value, bool):
                value = "True" if value else "False"

            cell = self._worksheet.cell(row=self._row_index, column=col_idx, value=value)

            if col_name in ("search_phrase_count", "contains_money", "date"):
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            if is_even:
                cell.fill = _ROW_FILL_EVEN

            cell.border = _THIN_BORDER

        self._worksheet.row_dimensions[self._row_index].height = 60
        self._row_index += 1
        logger.debug(f"Linha {self._row_index - 1} adicionada ao Excel.")

    def save(self) -> None:
        """Salva o workbook em disco com ajuste de colunas."""
        if self._workbook is None:
            logger.warning("Workbook nao inicializado. Nada a salvar.")
            return

        self.auto_fit_columns()
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self._workbook.save(self.output_path)
        total = self._row_index - 2
        logger.info(f"Excel salvo: {self.output_path} ({total} artigo(s))")

    def auto_fit_columns(self) -> None:
        """Ajusta a largura das colunas."""
        if self._worksheet is None:
            return
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            letter = get_column_letter(col_idx)
            self._worksheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(col_name, 20)
        logger.debug("Largura das colunas ajustada.")
